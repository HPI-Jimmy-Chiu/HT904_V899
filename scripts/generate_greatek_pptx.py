# -*- coding: utf-8 -*-
"""
產生 CC_Greatek 遠端控制操作說明書 PPTX
使用母版: HONPREC_PPT_templete_ver_C (SOP) .PPTX
輸出: RD5軟體_CC_Greatek_遠端控制操作說明書_20260415_170000.pptx
"""
import os
from pptx import Presentation
from pptx.util import Inches, Pt, Emu, Cm
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from copy import deepcopy
from lxml import etree
import datetime

# ── 路徑設定 ──────────────────────────────────────────────
TEMPLATE = r"D:\HT9045\document\HONPREC_PPT_templete_ver_C (SOP) .PPTX"
OUTPUT   = r"D:\HT9045\document\RD5軟體_CC_Greatek_遠端控制操作說明書_20260415_170000.pptx"

# ── 色彩常數 ──────────────────────────────────────────────
CLR_TITLE   = RGBColor(0x1A, 0x3C, 0x6E)   # 深藍
CLR_SUB     = RGBColor(0x33, 0x33, 0x33)   # 深灰
CLR_BODY    = RGBColor(0x33, 0x33, 0x33)   # 內文
CLR_WHITE   = RGBColor(0xFF, 0xFF, 0xFF)
CLR_ACCENT  = RGBColor(0xC0, 0x39, 0x2B)   # 紅色警示
CLR_GREEN   = RGBColor(0x27, 0xAE, 0x60)   # 綠色(成功)
CLR_ORANGE  = RGBColor(0xE6, 0x7E, 0x22)   # 橘色(警告)
CLR_BLUE    = RGBColor(0x29, 0x80, 0xB9)   # 藍色
CLR_TBL_HDR = RGBColor(0x1A, 0x3C, 0x6E)   # 表頭背景
CLR_TBL_ALT = RGBColor(0xEB, 0xF5, 0xFB)   # 表格交替行

# ── 字型 ──────────────────────────────────────────────────
FONT_MAIN = "微軟正黑體"
FONT_CODE = "Consolas"

# ── 開啟母版 ──────────────────────────────────────────────
prs = Presentation(TEMPLATE)

# 刪除所有範本投影片（保留母版）
while len(prs.slides) > 0:
    rId = prs.slides._sldIdLst[0].get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id')
    prs.part.drop_rel(rId)
    prs.slides._sldIdLst.remove(prs.slides._sldIdLst[0])

# ── 版面配置索引 ──────────────────────────────────────────
LO_COVER    = prs.slide_layouts[0]   # 標題投影片
LO_TOC      = prs.slide_layouts[3]   # 2_標題及物件 (4 組 title+text)
LO_CONTENT  = prs.slide_layouts[4]   # 1_標題及物件 (標題+副標+大內容區)
LO_BLANK    = prs.slide_layouts[2]   # 3_標題及物件 (空白可自由加形狀)
LO_END      = prs.slide_layouts[7]   # 自訂版面配置 (結尾)


# ── Helper Functions ──────────────────────────────────────

def add_textbox(slide, left, top, width, height):
    """新增文字方塊，回傳 text_frame"""
    from pptx.util import Emu
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    return tf

def set_para(tf, text, font_name=FONT_MAIN, font_size=Pt(14), bold=False,
             color=CLR_BODY, alignment=PP_ALIGN.LEFT, space_after=Pt(4),
             space_before=Pt(0), level=0):
    """在 text_frame 新增一個段落（不蓋掉第一段則需先呼叫一次再用 add_para）"""
    p = tf.paragraphs[-1] if tf.paragraphs[0].text == "" and len(tf.paragraphs) == 1 else tf.add_paragraph()
    p.text = text
    p.font.name = font_name
    p.font.size = font_size
    p.font.bold = bold
    p.font.color.rgb = color
    p.alignment = alignment
    p.space_after = space_after
    p.space_before = space_before
    p.level = level
    return p

def add_para(tf, text, font_name=FONT_MAIN, font_size=Pt(14), bold=False,
             color=CLR_BODY, alignment=PP_ALIGN.LEFT, space_after=Pt(4),
             space_before=Pt(0)):
    """追加段落"""
    p = tf.add_paragraph()
    p.text = text
    p.font.name = font_name
    p.font.size = font_size
    p.font.bold = bold
    p.font.color.rgb = color
    p.alignment = alignment
    p.space_after = space_after
    p.space_before = space_before
    return p

def add_table(slide, rows, cols, left, top, width, height):
    """新增表格"""
    table_shape = slide.shapes.add_table(rows, cols, left, top, width, height)
    return table_shape.table

def style_table_cell(cell, text, font_size=Pt(11), bold=False, color=CLR_BODY,
                     fill=None, alignment=PP_ALIGN.LEFT, font_name=FONT_MAIN):
    """設定表格儲存格"""
    cell.text = ""
    p = cell.text_frame.paragraphs[0]
    p.text = text
    p.font.name = font_name
    p.font.size = font_size
    p.font.bold = bold
    p.font.color.rgb = color
    p.alignment = alignment
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    if fill:
        cell.fill.solid()
        cell.fill.fore_color.rgb = fill

def make_content_slide(title_text, subtitle_text=""):
    """建立內容頁 (Layout 4: 1_標題及物件)"""
    slide = prs.slides.add_slide(LO_CONTENT)
    # 標題 (idx=0)
    if slide.placeholders[0]:
        ph = slide.placeholders[0]
        ph.text = title_text
        for p in ph.text_frame.paragraphs:
            p.font.name = FONT_MAIN
            p.font.size = Pt(24)
            p.font.bold = True
            p.font.color.rgb = CLR_TITLE
    return slide


# ══════════════════════════════════════════════════════════
#  Slide 1: 封面
# ══════════════════════════════════════════════════════════
slide = prs.slides.add_slide(LO_COVER)
# 標題
if 0 in [ph.placeholder_format.idx for ph in slide.placeholders]:
    ph = slide.placeholders[0]
    ph.text = "HT9045 Handler\n遠端控制操作說明書"
    for p in ph.text_frame.paragraphs:
        p.font.name = FONT_MAIN
        p.font.size = Pt(32)
        p.font.bold = True
        p.font.color.rgb = CLR_TITLE
        p.alignment = PP_ALIGN.LEFT

# 副標題 (idx=10)
if 10 in [ph.placeholder_format.idx for ph in slide.placeholders]:
    ph = slide.placeholders[10]
    ph.text = "TCP Socket 遠端控制功能 — 客戶端操作指南"
    for p in ph.text_frame.paragraphs:
        p.font.name = FONT_MAIN
        p.font.size = Pt(18)
        p.font.color.rgb = CLR_SUB

# 封面資訊欄位
info_fields = {
    11: "適用客戶：Greatek（超豐）",
    12: "文件版本：V1.0",
    14: "建立日期：2026-04-15",
    15: "文件性質：客戶端操作說明",
}
for idx, text in info_fields.items():
    try:
        ph = slide.placeholders[idx]
        ph.text = text
        for p in ph.text_frame.paragraphs:
            p.font.name = FONT_MAIN
            p.font.size = Pt(12)
            p.font.color.rgb = CLR_SUB
    except KeyError:
        pass


# ══════════════════════════════════════════════════════════
#  Slide 2: 目錄
# ══════════════════════════════════════════════════════════
slide = prs.slides.add_slide(LO_BLANK)

# 手動加標題
tf = add_textbox(slide, Cm(2.2), Cm(0.5), Cm(25), Cm(1.5))
set_para(tf, "目  錄", FONT_MAIN, Pt(28), True, CLR_TITLE, PP_ALIGN.LEFT)

toc_items = [
    ("1", "功能概述與實作狀態"),
    ("2", "Handler 端設定"),
    ("3", "客戶端 Socket 通訊規格"),
    ("4", "指令集說明"),
    ("5", "完整操作流程"),
    ("6", "錯誤處理與注意事項"),
    ("7", "附錄：快速參考表"),
]

tf2 = add_textbox(slide, Cm(3), Cm(3), Cm(20), Cm(12))
for num, label in toc_items:
    set_para(tf2, f"{num}.  {label}", FONT_MAIN, Pt(18), False, CLR_BODY, PP_ALIGN.LEFT, Pt(12), Pt(6))


# ══════════════════════════════════════════════════════════
#  Slide 3: 功能概述
# ══════════════════════════════════════════════════════════
slide = make_content_slide("1. 功能概述")

tf = add_textbox(slide, Cm(2.2), Cm(3), Cm(28), Cm(12))
set_para(tf, "HT9045 Handler 支援透過 TCP Socket 進行遠端控制，允許客戶端程式（Agent）執行：",
         FONT_MAIN, Pt(14), False, CLR_BODY)

features = [
    ("遠端啟動 (Auto Start)", "遠端觸發機台從閒置進入運行狀態"),
    ("清除批次 (End Lot)", "清除上一批次的統計數據與計數器"),
    ("啟動批次 (Start Lot)", "設定批次資訊（LotID、OperatorID、FT/RT）及下載 Setup File"),
    ("安全鎖定 / 解鎖", "遠端鎖定或解鎖機台，防止意外啟動"),
    ("機台狀態查詢", "查詢機台運行狀態、批次資訊、溫度等"),
]
for title, desc in features:
    add_para(tf, f"?  {title}", FONT_MAIN, Pt(14), True, CLR_BLUE, space_before=Pt(6))
    add_para(tf, f"     {desc}", FONT_MAIN, Pt(12), False, CLR_BODY)


# ══════════════════════════════════════════════════════════
#  Slide 4: 功能實作狀態
# ══════════════════════════════════════════════════════════
slide = make_content_slide("功能實作狀態")

tbl = add_table(slide, 7, 4, Cm(2.2), Cm(3.5), Cm(28), Cm(10))
headers = ["功能", "規格書需求", "實作狀態", "備註"]
data = [
    ["Auto Start（遠端啟動）",      "Requirement #1", "? 已完成", "機台需處於 HALT"],
    ["Safety Interlock（Lock/Unlock）", "Requirement #3", "? 已完成", "Lock 時無法啟動"],
    ["End Lot（清除批次）",          "Requirement #5", "? 已完成", "支援全清及分項"],
    ["Start Lot — 批次資訊",       "Requirement #4-1", "? 已完成", "LotID / OperatorID / FT|RT"],
    ["Start Lot — Setup File",     "Requirement #4-2", "? 已完成", "zip 壓縮格式"],
    ["One Cycle 模式",              "Requirement #2", "?? 僅除錯模式", "正式版不支援"],
]

col_widths = [Cm(8), Cm(5.5), Cm(5), Cm(9.5)]
for i, w in enumerate(col_widths):
    tbl.columns[i].width = w

for j, h in enumerate(headers):
    style_table_cell(tbl.cell(0, j), h, Pt(12), True, CLR_WHITE, CLR_TBL_HDR, PP_ALIGN.CENTER)

for i, row in enumerate(data):
    for j, val in enumerate(row):
        fill = CLR_TBL_ALT if i % 2 == 0 else None
        style_table_cell(tbl.cell(i+1, j), val, Pt(11), False, CLR_BODY, fill)


# ══════════════════════════════════════════════════════════
#  Slide 5: Handler 端設定
# ══════════════════════════════════════════════════════════
slide = make_content_slide("2. Handler 端設定")

tf = add_textbox(slide, Cm(2.2), Cm(3), Cm(28), Cm(14))
set_para(tf, "2.1 客戶碼確認", FONT_MAIN, Pt(18), True, CLR_TITLE, space_before=Pt(4))
add_para(tf, "Handler 需設定為 Greatek（超豐）客戶碼，系統啟動時自動開啟 OLP Server。", space_before=Pt(4))
add_para(tf, "此設定由設備供應商在出廠或維護時配置，客戶端無需調整。", FONT_MAIN, Pt(12), False, CLR_SUB)

add_para(tf, "", space_after=Pt(4))
add_para(tf, "2.2 通訊服務狀態確認", FONT_MAIN, Pt(18), True, CLR_TITLE, space_before=Pt(8))
add_para(tf, "?  Handler 啟動後，在 Automation 監控畫面確認 OLP 狀態已開啟", space_before=Pt(4))
add_para(tf, "?  客戶端成功連線時，畫面會顯示連線狀態")

add_para(tf, "", space_after=Pt(4))
add_para(tf, "2.3 網路環境要求", FONT_MAIN, Pt(18), True, CLR_TITLE, space_before=Pt(8))
add_para(tf, "?  Handler PC 需具備乙太網路連線", space_before=Pt(4))
add_para(tf, "?  防火牆需開放 TCP Port 6671（入站規則）", FONT_MAIN, Pt(14), True, CLR_ACCENT)
add_para(tf, "?  客戶端 PC 與 Handler PC 需在同一網段或可路由網段")


# ══════════════════════════════════════════════════════════
#  Slide 6: Socket 通訊規格 — 連線參數
# ══════════════════════════════════════════════════════════
slide = make_content_slide("3. 客戶端 Socket 通訊規格")

tbl = add_table(slide, 7, 2, Cm(2.2), Cm(3.5), Cm(16), Cm(8.5))
conn_params = [
    ["通訊協定",   "TCP/IP"],
    ["連線模式",   "Client → Server（客戶端主動連線 Handler）"],
    ["Handler IP", "Handler PC 的 IP 位址（請洽設備管理人員）"],
    ["Port",       "6671"],
    ["編碼",       "ASCII"],
    ["最大連線數", "1（同一時間僅允許一個客戶端連線）"],
]
style_table_cell(tbl.cell(0, 0), "參數", Pt(12), True, CLR_WHITE, CLR_TBL_HDR, PP_ALIGN.CENTER)
style_table_cell(tbl.cell(0, 1), "值",   Pt(12), True, CLR_WHITE, CLR_TBL_HDR, PP_ALIGN.CENTER)
tbl.columns[0].width = Cm(5)
tbl.columns[1].width = Cm(11)
for i, (k, v) in enumerate(conn_params):
    fill = CLR_TBL_ALT if i % 2 == 0 else None
    style_table_cell(tbl.cell(i+1, 0), k, Pt(11), True, CLR_BODY, fill)
    style_table_cell(tbl.cell(i+1, 1), v, Pt(11), False, CLR_BODY, fill)

# 警告
tf = add_textbox(slide, Cm(2.2), Cm(13), Cm(28), Cm(2))
set_para(tf, "??  若連線數超過上限，Handler 將自動重置 OLP Server。請確保只有一個客戶端連線。",
         FONT_MAIN, Pt(12), True, CLR_ACCENT)


# ══════════════════════════════════════════════════════════
#  Slide 7: 封包格式
# ══════════════════════════════════════════════════════════
slide = make_content_slide("封包格式")

tf = add_textbox(slide, Cm(2.2), Cm(3), Cm(28), Cm(14))
set_para(tf, "所有指令與回傳均採用以下固定封包格式：", FONT_MAIN, Pt(14))

add_para(tf, "", space_after=Pt(2))
add_para(tf, "  [STX] + Command + [SOH] + V_Total + [SOH] + Data[0] + ... + [ETX]",
         FONT_CODE, Pt(14), True, CLR_BLUE, space_before=Pt(4))

add_para(tf, "", space_after=Pt(8))
add_para(tf, "控制字元定義：", FONT_MAIN, Pt(14), True, CLR_TITLE, space_before=Pt(8))

ctrl_chars = [
    ("STX", "0x02", "封包起始（Start of Text）"),
    ("SOH", "0x01", "欄位分隔符（Start of Heading）"),
    ("ETX", "0x03", "封包結束（End of Text）"),
]
for sym, hex_val, desc in ctrl_chars:
    add_para(tf, f"  {sym}  =  {hex_val}    {desc}", FONT_CODE, Pt(12), False, CLR_BODY, space_before=Pt(2))

add_para(tf, "", space_after=Pt(8))
add_para(tf, "封包組裝範例：", FONT_MAIN, Pt(14), True, CLR_TITLE, space_before=Pt(8))
add_para(tf, '  Auto Start:  0x02 "START_REQUEST" 0x01 "0" 0x03', FONT_CODE, Pt(12), False, CLR_BODY, space_before=Pt(4))
add_para(tf, '  End Lot:     0x02 "CLEAR_REPORT_REQUEST" 0x01 "1" 0x01 "1111" 0x03', FONT_CODE, Pt(12))
add_para(tf, '  Start Lot:   0x02 "LotInfo_REQUEST" 0x01 "3" 0x01 "LOT001" 0x01 "OP001" 0x01 "0" 0x03', FONT_CODE, Pt(12))


# ══════════════════════════════════════════════════════════
#  Slide 8: 指令集 — Safety Unlock / Lock
# ══════════════════════════════════════════════════════════
slide = make_content_slide("4. 指令集 — Safety Unlock / Lock")

tf = add_textbox(slide, Cm(2.2), Cm(3), Cm(28), Cm(14))
set_para(tf, "Safety Unlock — 解除安全鎖定", FONT_MAIN, Pt(16), True, CLR_GREEN, space_before=Pt(4))
add_para(tf, "啟動機台前，必須先解除安全鎖定。", space_before=Pt(4))
add_para(tf, "  發送：RESUME_REQUEST   V_Total=0", FONT_CODE, Pt(12), False, CLR_BODY, space_before=Pt(4))
add_para(tf, "  回傳：RESUME_REPLY     Data[0]=0（成功）", FONT_CODE, Pt(12))

add_para(tf, "", space_after=Pt(8))
add_para(tf, "Safety Lock — 安全鎖定", FONT_MAIN, Pt(16), True, CLR_ACCENT, space_before=Pt(8))
add_para(tf, "良率異常或需緊急停止時，鎖定機台。鎖定狀態下 Auto Start 將被拒絕。", space_before=Pt(4))
add_para(tf, "  發送：PAUSE_REQUEST    V_Total=0", FONT_CODE, Pt(12), False, CLR_BODY, space_before=Pt(4))
add_para(tf, "  回傳：PAUSE_REPLY      Data[0]=0（成功）", FONT_CODE, Pt(12))
add_para(tf, "", space_after=Pt(4))
add_para(tf, "??  Lock 效果立即生效，運行中機台將觸發 Soft Stop", FONT_MAIN, Pt(12), True, CLR_ACCENT, space_before=Pt(4))


# ══════════════════════════════════════════════════════════
#  Slide 9: 指令集 — End Lot
# ══════════════════════════════════════════════════════════
slide = make_content_slide("指令集 — End Lot（清除批次）")

tf = add_textbox(slide, Cm(2.2), Cm(3), Cm(28), Cm(2.5))
set_para(tf, "清除上一批次的統計數據，為新批次做準備。", FONT_MAIN, Pt(14))
add_para(tf, "前提條件：機台必須處於停止狀態（非運行中）", FONT_MAIN, Pt(12), True, CLR_ACCENT, space_before=Pt(4))

# 發送格式
tf2 = add_textbox(slide, Cm(2.2), Cm(5.8), Cm(28), Cm(2))
set_para(tf2, "  發送：CLEAR_REPORT_REQUEST   V_Total=1   Data[0]=清除代碼", FONT_CODE, Pt(12), False, CLR_BODY)
add_para(tf2, "  回傳：CLEAR_REPORT_REPLY     Data[0]=0（成功）", FONT_CODE, Pt(12))

# 清除代碼表
tbl = add_table(slide, 7, 2, Cm(2.2), Cm(8.5), Cm(16), Cm(7))
style_table_cell(tbl.cell(0, 0), "代碼",   Pt(12), True, CLR_WHITE, CLR_TBL_HDR, PP_ALIGN.CENTER)
style_table_cell(tbl.cell(0, 1), "清除範圍", Pt(12), True, CLR_WHITE, CLR_TBL_HDR, PP_ALIGN.CENTER)
tbl.columns[0].width = Cm(3)
tbl.columns[1].width = Cm(13)
codes = [
    ("1111", "全部清除（推薦用於批次切換）"),
    ("2",    "僅清除時間資料"),
    ("3",    "僅清除裝載計數"),
    ("4",    "僅清除 Tray Sort 計數"),
    ("5",    "僅清除 Contact 計數"),
    ("6",    "僅清除測試分類計數"),
]
for i, (code, desc) in enumerate(codes):
    fill = CLR_TBL_ALT if i % 2 == 0 else None
    bld = True if code == "1111" else False
    style_table_cell(tbl.cell(i+1, 0), code, Pt(11), bld, CLR_BODY, fill, PP_ALIGN.CENTER, FONT_CODE)
    style_table_cell(tbl.cell(i+1, 1), desc, Pt(11), bld, CLR_BODY, fill)


# ══════════════════════════════════════════════════════════
#  Slide 10: 指令集 — Start Lot (批次資訊)
# ══════════════════════════════════════════════════════════
slide = make_content_slide("指令集 — Start Lot（批次資訊設定）")

tf = add_textbox(slide, Cm(2.2), Cm(3), Cm(28), Cm(14))
set_para(tf, "設定新批次的 Lot ID、操作員 ID 及測試模式。", FONT_MAIN, Pt(14))
add_para(tf, "前提條件：機台非運行中，且內部無殘留 IC。", FONT_MAIN, Pt(12), True, CLR_ACCENT, space_before=Pt(4))

add_para(tf, "", space_after=Pt(6))
add_para(tf, "發送格式：", FONT_MAIN, Pt(14), True, CLR_TITLE, space_before=Pt(6))
add_para(tf, "  Command  = LotInfo_REQUEST", FONT_CODE, Pt(12), False, CLR_BODY, space_before=Pt(2))
add_para(tf, "  V_Total  = 3", FONT_CODE, Pt(12))
add_para(tf, "  Data[0]  = Lot ID（字串）", FONT_CODE, Pt(12))
add_para(tf, "  Data[1]  = Operator ID（字串）", FONT_CODE, Pt(12))
add_para(tf, "  Data[2]  = 0（FT）或 1（RT）", FONT_CODE, Pt(12))

add_para(tf, "", space_after=Pt(6))
add_para(tf, "回傳格式：", FONT_MAIN, Pt(14), True, CLR_TITLE, space_before=Pt(6))
add_para(tf, "  Command  = LotInfo_REPLY", FONT_CODE, Pt(12), False, CLR_BODY, space_before=Pt(2))
add_para(tf, "  Data[0]  = 0（成功）", FONT_CODE, Pt(12), False, CLR_GREEN)
add_para(tf, "  Data[0]  = 2（失敗：機台運行中或有 IC 在機台內）", FONT_CODE, Pt(12), False, CLR_ACCENT)


# ══════════════════════════════════════════════════════════
#  Slide 11: 指令集 — Start Lot (Setup File)
# ══════════════════════════════════════════════════════════
slide = make_content_slide("指令集 — Start Lot（Setup File 下載）")

tf = add_textbox(slide, Cm(2.2), Cm(3), Cm(28), Cm(14))
set_para(tf, "將 Setup File（zip 壓縮格式）傳送至 Handler 並自動解壓套用。", FONT_MAIN, Pt(14))

add_para(tf, "", space_after=Pt(4))
add_para(tf, "發送格式：", FONT_MAIN, Pt(14), True, CLR_TITLE, space_before=Pt(6))
add_para(tf, "  Command  = PP_DL_REQUEST", FONT_CODE, Pt(12), False, CLR_BODY, space_before=Pt(2))
add_para(tf, "  V_Total  = 3", FONT_CODE, Pt(12))
add_para(tf, "  Data[0]  = 檔案名稱（含 .zip 副檔名）", FONT_CODE, Pt(12))
add_para(tf, "  Data[1]  = 檔案大小（byte，字串）", FONT_CODE, Pt(12))
add_para(tf, "  Data[2]  = 檔案內容（Hex 字串編碼，每 byte 以 2 碼 hex 表示）", FONT_CODE, Pt(12))

add_para(tf, "", space_after=Pt(4))
add_para(tf, "檔案編碼方式：", FONT_MAIN, Pt(14), True, CLR_TITLE, space_before=Pt(6))
add_para(tf, '  將 zip 檔案的每個 byte 轉換為 2 碼 hex ASCII 字串', FONT_MAIN, Pt(12), space_before=Pt(2))
add_para(tf, '  例：byte 0xAB → 字串 "AB"', FONT_CODE, Pt(12), False, CLR_BLUE)

add_para(tf, "", space_after=Pt(4))
add_para(tf, "回傳格式：", FONT_MAIN, Pt(14), True, CLR_TITLE, space_before=Pt(6))
add_para(tf, "  Command  = PP_DL_REPLY", FONT_CODE, Pt(12), False, CLR_BODY, space_before=Pt(2))
add_para(tf, "  Data[0]  = 1（成功）", FONT_CODE, Pt(12), False, CLR_GREEN)
add_para(tf, "  Data[0]  = 0（失敗）", FONT_CODE, Pt(12), False, CLR_ACCENT)


# ══════════════════════════════════════════════════════════
#  Slide 12: 指令集 — Auto Start
# ══════════════════════════════════════════════════════════
slide = make_content_slide("指令集 — Auto Start（遠端啟動）")

tf = add_textbox(slide, Cm(2.2), Cm(3), Cm(28), Cm(5))
set_para(tf, "觸發機台從閒置狀態進入運行狀態。", FONT_MAIN, Pt(14))
add_para(tf, "", space_after=Pt(2))
add_para(tf, "  發送：START_REQUEST      V_Total=0", FONT_CODE, Pt(13), False, CLR_BODY, space_before=Pt(4))
add_para(tf, "  回傳：START_REPLY        Data[0]=0（固定回傳）", FONT_CODE, Pt(13))

# 執行條件表
tbl = add_table(slide, 4, 2, Cm(2.2), Cm(7.5), Cm(20), Cm(4.5))
style_table_cell(tbl.cell(0, 0), "執行條件", Pt(12), True, CLR_WHITE, CLR_TBL_HDR, PP_ALIGN.CENTER)
style_table_cell(tbl.cell(0, 1), "說明",     Pt(12), True, CLR_WHITE, CLR_TBL_HDR, PP_ALIGN.CENTER)
tbl.columns[0].width = Cm(7)
tbl.columns[1].width = Cm(13)
conds = [
    ("機台狀態 = HALT", "機台必須處於閒置停止狀態"),
    ("未運行中", "機台尚未啟動"),
    ("未被鎖定", "安全鎖定必須已解除（需先送 RESUME_REQUEST）"),
]
for i, (k, v) in enumerate(conds):
    fill = CLR_TBL_ALT if i % 2 == 0 else None
    style_table_cell(tbl.cell(i+1, 0), k, Pt(11), True, CLR_BODY, fill)
    style_table_cell(tbl.cell(i+1, 1), v, Pt(11), False, CLR_BODY, fill)

tf2 = add_textbox(slide, Cm(2.2), Cm(12.5), Cm(28), Cm(2))
set_para(tf2, "??  若條件不滿足，Handler 仍回傳 START_REPLY，但不會實際啟動機台。",
         FONT_MAIN, Pt(12), True, CLR_ACCENT)
add_para(tf2, "建議先送 MAIN_STATUS_INQUIRE 確認狀態為 HALT 再送 Auto Start。",
         FONT_MAIN, Pt(12), False, CLR_SUB)


# ══════════════════════════════════════════════════════════
#  Slide 13: 指令集 — 狀態查詢
# ══════════════════════════════════════════════════════════
slide = make_content_slide("指令集 — 狀態查詢（INQUIRE）")

# 機台狀態
tf = add_textbox(slide, Cm(2.2), Cm(3), Cm(28), Cm(5))
set_para(tf, "機台狀態查詢", FONT_MAIN, Pt(16), True, CLR_TITLE)
add_para(tf, "  發送：MAIN_STATUS_INQUIRE   V_Total=0", FONT_CODE, Pt(12), False, CLR_BODY, space_before=Pt(4))
add_para(tf, "  回傳：MAIN_STATUS_GRANT     Data[0]=狀態字串", FONT_CODE, Pt(12))

tbl = add_table(slide, 5, 2, Cm(2.2), Cm(5.8), Cm(14), Cm(4.8))
style_table_cell(tbl.cell(0, 0), "狀態值", Pt(11), True, CLR_WHITE, CLR_TBL_HDR, PP_ALIGN.CENTER)
style_table_cell(tbl.cell(0, 1), "說明",   Pt(11), True, CLR_WHITE, CLR_TBL_HDR, PP_ALIGN.CENTER)
tbl.columns[0].width = Cm(4)
tbl.columns[1].width = Cm(10)
statuses = [
    ("RUN",   "運行中"),
    ("PAUSE", "暫停中"),
    ("HALT",  "停止（可接受 Auto Start）"),
    ("LOCK",  "鎖定中（需先解鎖）"),
]
for i, (s, d) in enumerate(statuses):
    fill = CLR_TBL_ALT if i % 2 == 0 else None
    style_table_cell(tbl.cell(i+1, 0), s, Pt(11), True, CLR_BODY, fill, PP_ALIGN.CENTER, FONT_CODE)
    style_table_cell(tbl.cell(i+1, 1), d, Pt(11), False, CLR_BODY, fill)

# 其他查詢
tf2 = add_textbox(slide, Cm(2.2), Cm(11.2), Cm(28), Cm(5))
set_para(tf2, "其他查詢指令：", FONT_MAIN, Pt(14), True, CLR_TITLE, space_before=Pt(4))
add_para(tf2, "  LotInfo_INQUIRE          → LotInfo_GRANT（批次資訊）", FONT_CODE, Pt(12), space_before=Pt(4))
add_para(tf2, "  SETUP_FILE_NAME_INQUIRE  → SETUP_FILE_NAME_GRANT（Setup File 名稱）", FONT_CODE, Pt(12))
add_para(tf2, "  VERSION_INQUIRE          → VERSION_GRANT（軟體版本）", FONT_CODE, Pt(12))


# ══════════════════════════════════════════════════════════
#  Slide 14: 完整操作流程 — Auto Start
# ══════════════════════════════════════════════════════════
slide = make_content_slide("5. 完整操作流程 — Auto Start")

tf = add_textbox(slide, Cm(2.2), Cm(3), Cm(28), Cm(14))
set_para(tf, "標準 Auto Start 流程（依序執行）：", FONT_MAIN, Pt(14), True, CLR_TITLE)

steps = [
    ("步驟 1", "建立 TCP 連線", "Client 連線至 Handler IP:6671"),
    ("步驟 2", "查詢機台狀態（建議）", "MAIN_STATUS_INQUIRE → 確認為 HALT"),
    ("步驟 3", "清除上一批次 (End Lot)", "CLEAR_REPORT_REQUEST [1111]"),
    ("步驟 4", "設定批次資訊", "LotInfo_REQUEST [LotID, OperatorID, 0/1] → 確認回傳碼 = 0"),
    ("步驟 5", "下載 Setup File（如需要）", "PP_DL_REQUEST [filename.zip, size, hex] → 確認回傳碼 = 1"),
    ("步驟 6", "解除安全鎖定", "RESUME_REQUEST → RESUME_REPLY"),
    ("步驟 7", "遠端啟動 (Auto Start)", "START_REQUEST → START_REPLY"),
    ("步驟 8", "確認機台已啟動（建議）", "MAIN_STATUS_INQUIRE → 確認為 RUN"),
]
for step, title, detail in steps:
    add_para(tf, f"  {step}：{title}", FONT_MAIN, Pt(13), True, CLR_BLUE, space_before=Pt(6))
    add_para(tf, f"           {detail}", FONT_CODE, Pt(11), False, CLR_BODY, space_before=Pt(1))


# ══════════════════════════════════════════════════════════
#  Slide 15: 緊急鎖定 & 重新啟動
# ══════════════════════════════════════════════════════════
slide = make_content_slide("緊急鎖定 / 解除後重新啟動")

tf = add_textbox(slide, Cm(2.2), Cm(3), Cm(28), Cm(14))
set_para(tf, "緊急鎖定流程", FONT_MAIN, Pt(18), True, CLR_ACCENT)
add_para(tf, "當偵測到異常（如良率異常）時：", space_before=Pt(6))
add_para(tf, "  1. 發送 PAUSE_REQUEST → Handler 立即鎖定", FONT_CODE, Pt(13), False, CLR_BODY, space_before=Pt(4))
add_para(tf, "  2. 運行中的機台將執行 Soft Stop", FONT_CODE, Pt(13))
add_para(tf, "  3. 等待 PAUSE_REPLY", FONT_CODE, Pt(13))

add_para(tf, "", space_after=Pt(10))
add_para(tf, "解除鎖定後重新啟動", FONT_MAIN, Pt(18), True, CLR_GREEN, space_before=Pt(10))
add_para(tf, "  1. 發送 RESUME_REQUEST → 解除鎖定", FONT_CODE, Pt(13), False, CLR_BODY, space_before=Pt(6))
add_para(tf, "  2. 查詢 MAIN_STATUS_INQUIRE → 確認為 HALT", FONT_CODE, Pt(13))
add_para(tf, "  3. 發送 START_REQUEST → 重新啟動", FONT_CODE, Pt(13))


# ══════════════════════════════════════════════════════════
#  Slide 16: 錯誤處理
# ══════════════════════════════════════════════════════════
slide = make_content_slide("6. 錯誤處理與注意事項")

tbl = add_table(slide, 6, 3, Cm(2.2), Cm(3.5), Cm(28), Cm(7))
style_table_cell(tbl.cell(0, 0), "情境",     Pt(11), True, CLR_WHITE, CLR_TBL_HDR, PP_ALIGN.CENTER)
style_table_cell(tbl.cell(0, 1), "現象",     Pt(11), True, CLR_WHITE, CLR_TBL_HDR, PP_ALIGN.CENTER)
style_table_cell(tbl.cell(0, 2), "處理建議", Pt(11), True, CLR_WHITE, CLR_TBL_HDR, PP_ALIGN.CENTER)
tbl.columns[0].width = Cm(9)
tbl.columns[1].width = Cm(8)
tbl.columns[2].width = Cm(11)
errors = [
    ("非 HALT 狀態送 Auto Start", "不執行啟動，仍回傳 REPLY", "先查詢確認為 HALT"),
    ("機台運行中送 LotInfo", "回傳碼 = 2（失敗）", "等待停止後再設定"),
    ("Setup File 解壓失敗", "PP_DL_REPLY 回傳 0", "檢查 zip 完整性與格式"),
    ("連線中斷", "無回傳", "重新建立 TCP 連線"),
    ("安全鎖定未解除送 Auto Start", "機台不啟動", "先送 RESUME_REQUEST"),
]
for i, (s, p, a) in enumerate(errors):
    fill = CLR_TBL_ALT if i % 2 == 0 else None
    style_table_cell(tbl.cell(i+1, 0), s, Pt(10), False, CLR_BODY, fill)
    style_table_cell(tbl.cell(i+1, 1), p, Pt(10), False, CLR_BODY, fill)
    style_table_cell(tbl.cell(i+1, 2), a, Pt(10), False, CLR_BODY, fill)

# 逾時建議
tbl2 = add_table(slide, 5, 2, Cm(2.2), Cm(11.5), Cm(16), Cm(4.5))
style_table_cell(tbl2.cell(0, 0), "指令", Pt(11), True, CLR_WHITE, CLR_TBL_HDR, PP_ALIGN.CENTER)
style_table_cell(tbl2.cell(0, 1), "建議逾時", Pt(11), True, CLR_WHITE, CLR_TBL_HDR, PP_ALIGN.CENTER)
tbl2.columns[0].width = Cm(10)
tbl2.columns[1].width = Cm(6)
timeouts = [
    ("一般指令（Start/Pause/Resume/Clear）", "5 秒"),
    ("LotInfo_REQUEST", "5 秒"),
    ("PP_DL_REQUEST（Setup File）", "30 秒"),
    ("狀態查詢（INQUIRE）", "3 秒"),
]
for i, (cmd, t) in enumerate(timeouts):
    fill = CLR_TBL_ALT if i % 2 == 0 else None
    style_table_cell(tbl2.cell(i+1, 0), cmd, Pt(10), False, CLR_BODY, fill)
    style_table_cell(tbl2.cell(i+1, 1), t,   Pt(10), True, CLR_BODY, fill, PP_ALIGN.CENTER)


# ══════════════════════════════════════════════════════════
#  Slide 17: 注意事項
# ══════════════════════════════════════════════════════════
slide = make_content_slide("重要注意事項")

tf = add_textbox(slide, Cm(2.2), Cm(3), Cm(28), Cm(14))
notes = [
    "同一時間僅允許一個客戶端連線，多重連線將導致 Server 重置",
    "指令需逐一發送，等待前一個指令回傳後再發送下一個",
    "Safety Lock (PAUSE_REQUEST) 效果立即生效，運行中機台將觸發 Soft Stop",
    "Setup File 必須為 zip 格式，檔案內容以 Hex 字串傳送",
    "批次切換時建議使用清除代碼 1111（全部清除），確保上一批資料完全清空",
]
for i, note in enumerate(notes):
    set_para(tf, f"  {i+1}.  {note}", FONT_MAIN, Pt(14), False, CLR_BODY, space_before=Pt(10))


# ══════════════════════════════════════════════════════════
#  Slide 18: 附錄 — 指令速查表
# ══════════════════════════════════════════════════════════
slide = make_content_slide("7. 附錄 — 指令速查表")

tbl = add_table(slide, 11, 3, Cm(1.5), Cm(3.5), Cm(30), Cm(12))
style_table_cell(tbl.cell(0, 0), "功能",           Pt(11), True, CLR_WHITE, CLR_TBL_HDR, PP_ALIGN.CENTER)
style_table_cell(tbl.cell(0, 1), "指令（Request）", Pt(11), True, CLR_WHITE, CLR_TBL_HDR, PP_ALIGN.CENTER)
style_table_cell(tbl.cell(0, 2), "回傳（Reply）",   Pt(11), True, CLR_WHITE, CLR_TBL_HDR, PP_ALIGN.CENTER)
tbl.columns[0].width = Cm(7)
tbl.columns[1].width = Cm(12)
tbl.columns[2].width = Cm(11)

cmds = [
    ("安全解鎖",      "RESUME_REQUEST",            "RESUME_REPLY"),
    ("安全鎖定",      "PAUSE_REQUEST",             "PAUSE_REPLY"),
    ("清除批次",      "CLEAR_REPORT_REQUEST",      "CLEAR_REPORT_REPLY"),
    ("設定批次",      "LotInfo_REQUEST",           "LotInfo_REPLY"),
    ("下載 Setup File", "PP_DL_REQUEST",           "PP_DL_REPLY"),
    ("遠端啟動",      "START_REQUEST",             "START_REPLY"),
    ("查詢狀態",      "MAIN_STATUS_INQUIRE",       "MAIN_STATUS_GRANT"),
    ("查詢批次",      "LotInfo_INQUIRE",           "LotInfo_GRANT"),
    ("查詢版本",      "VERSION_INQUIRE",           "VERSION_GRANT"),
    ("查詢 Setup File", "SETUP_FILE_NAME_INQUIRE", "SETUP_FILE_NAME_GRANT"),
]
for i, (func, req, rep) in enumerate(cmds):
    fill = CLR_TBL_ALT if i % 2 == 0 else None
    style_table_cell(tbl.cell(i+1, 0), func, Pt(10), False, CLR_BODY, fill)
    style_table_cell(tbl.cell(i+1, 1), req,  Pt(10), False, CLR_BODY, fill, PP_ALIGN.LEFT, FONT_CODE)
    style_table_cell(tbl.cell(i+1, 2), rep,  Pt(10), False, CLR_BODY, fill, PP_ALIGN.LEFT, FONT_CODE)


# ══════════════════════════════════════════════════════════
#  Slide 19: 附錄 — TCP 連線速查
# ══════════════════════════════════════════════════════════
slide = make_content_slide("附錄 — TCP 連線速查")

tf = add_textbox(slide, Cm(4), Cm(4), Cm(24), Cm(10))
lines = [
    ("Protocol", "TCP/IP"),
    ("IP",       "[Handler PC IP]"),
    ("Port",     "6671"),
    ("Mode",     "Client connects to Handler (Server)"),
    ("Encoding", "ASCII"),
    ("Max Conn", "1"),
]
for label, val in lines:
    set_para(tf, f"  {label:<12}:  {val}", FONT_CODE, Pt(16), False, CLR_BODY, space_before=Pt(6))

add_para(tf, "", space_after=Pt(12))
add_para(tf, "封包格式：", FONT_MAIN, Pt(16), True, CLR_TITLE, space_before=Pt(12))
add_para(tf, "  STX(0x02) + Command + SOH(0x01) + V_Total + [SOH(0x01) + Data]... + ETX(0x03)",
         FONT_CODE, Pt(13), True, CLR_BLUE, space_before=Pt(6))


# ══════════════════════════════════════════════════════════
#  Slide 20: 結尾頁
# ══════════════════════════════════════════════════════════
slide = prs.slides.add_slide(LO_END)
if 0 in [ph.placeholder_format.idx for ph in slide.placeholders]:
    ph = slide.placeholders[0]
    ph.text = "Thank You"
    for p in ph.text_frame.paragraphs:
        p.font.name = FONT_MAIN
        p.font.size = Pt(36)
        p.font.bold = True
        p.font.color.rgb = CLR_TITLE
        p.alignment = PP_ALIGN.CENTER


# ── 儲存 ─────────────────────────────────────────────────
prs.save(OUTPUT)
print(f"PPTX saved to: {OUTPUT}")
print(f"Total slides: {len(prs.slides)}")
